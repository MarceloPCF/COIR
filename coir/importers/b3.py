"""Leitura das planilhas da Área do Investidor da B3 (xlsx): movimentação, negociação e o
relatório consolidado mensal (posição).

Cuidados do formato (exportações de 2019-11 a 2026-09):
- o xlsx declara dimensão errada (1x1): ``reset_dimensions`` antes de ler;
- datas vêm como texto DD/MM/AAAA; "-" e vazio são "sem valor";
- números vêm como número (float do xlsx, convertido com ``str`` para não herdar a imprecisão) ou
  como texto com ponto de milhar ("1.000");
- a planilha consolidada termina com uma linha vazia e uma linha "Total";
- o consolidado mensal pode repetir séries de opção com Produto " - Opção" (sem o nome);
- cabeçalho desconhecido falha alto (a B3 mudou o layout).
"""

from collections import Counter
from collections.abc import Iterator
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Literal

from openpyxl import load_workbook

from coir.domain.b3 import LinhaMovimentacao, LinhaNegociacao, LinhaPosicao, TipoPlanilha
from coir.erros import ImportacaoError

ABA_MOVIMENTACAO = "Movimentação"
ABA_NEGOCIACAO = "Negociação"
ABA_POSICAO_ACOES = "Posição - Ações"
ABA_POSICAO_OPCOES = "Posição - Opções"
ABA_POSICAO_EMPRESTIMOS = "Posição - Empréstimos"
ABA_POSICAO_FUTURO = "Posição - Futuro"
MESES = ("janeiro", "fevereiro", "marco", "abril", "maio", "junho", "julho", "agosto", "setembro",
         "outubro", "novembro", "dezembro")  # fmt: skip
COLUNAS = {
    TipoPlanilha.MOVIMENTACAO: ("Entrada/Saída", "Data", "Movimentação", "Produto", "Instituição",
                                "Quantidade", "Preço unitário", "Valor da Operação"),
    TipoPlanilha.NEGOCIACAO: ("Data do Negócio", "Tipo de Movimentação", "Mercado",
                              "Prazo/Vencimento", "Instituição", "Código de Negociação",
                              "Quantidade", "Preço", "Valor"),
}  # fmt: skip
COLUNAS_ACOES = ("Produto", "Instituição", "Conta", "Código de Negociação",
                 "Código ISIN / Distribuição", "Quantidade")  # fmt: skip
COLUNAS_OPCOES = ("Produto", "Instituição", "Código de Negociação", "Condição", "Quantidade")
COLUNAS_EMPRESTIMOS = ("Produto", "Instituição", "Natureza", "Quantidade")
COLUNAS_FUTURO = ("Produto", "Instituição", "Condição", "Quantidade")


def tipo(caminho: Path) -> TipoPlanilha:
    """Tipo pela aba (o nome do arquivo muda conforme a data da exportação)."""
    abas = set(_abas(caminho))
    if ABA_MOVIMENTACAO in abas:
        return TipoPlanilha.MOVIMENTACAO
    if ABA_NEGOCIACAO in abas:
        return TipoPlanilha.NEGOCIACAO
    if any(a.startswith("Posição - ") for a in abas):
        return TipoPlanilha.POSICAO
    raise ImportacaoError(f"{caminho.name}: não é planilha de movimentação, negociação ou "
                          f"consolidado da B3 (abas: {', '.join(sorted(abas))})")  # fmt: skip


def _abas(caminho: Path) -> list[str]:
    livro = load_workbook(caminho, read_only=True)
    try:
        return list(livro.sheetnames)
    finally:
        livro.close()


def _linhas(caminho: Path, aba: str, exigidas: tuple[str, ...]) -> Iterator[dict[str, Any]]:
    livro = load_workbook(caminho, read_only=True, data_only=True)
    try:
        folha = livro[aba]
        folha.reset_dimensions()  # type: ignore[union-attr]
        linhas = folha.iter_rows(values_only=True)  # type: ignore[union-attr]
        cabecalho = [str(c).strip() if c is not None else "" for c in next(linhas, ())]
        faltando = [c for c in exigidas if c not in cabecalho]
        if faltando:
            raise ImportacaoError(f"{caminho.name}, aba {aba}: colunas ausentes "
                                  f"{', '.join(faltando)} (a B3 mudou o layout?)")  # fmt: skip
        for valores in linhas:
            if valores is None or all(v in (None, "") for v in valores):
                continue
            linha = dict(zip(cabecalho, valores, strict=False))
            if _vazio(linha.get(exigidas[0])) or any(
                str(v).strip() == "Total" for v in valores if v is not None
            ):
                continue  # rodapé do consolidado (linha "Total" ou só com o valor somado)
            yield linha
    finally:
        livro.close()


def _texto(v: Any) -> str:
    return "" if v is None else str(v).strip()


def _vazio(v: Any) -> bool:
    return v is None or _texto(v) in ("", "-")


def _data(v: Any, contexto: str) -> date:
    if isinstance(v, datetime):
        return v.date()
    try:
        return datetime.strptime(_texto(v), "%d/%m/%Y").date()
    except ValueError as erro:
        raise ImportacaoError(f"{contexto}: data ilegível {v!r}") from erro


def _numero(v: Any, contexto: str) -> Decimal | None:
    if _vazio(v):
        return None
    if isinstance(v, bool):
        raise ImportacaoError(f"{contexto}: número ilegível {v!r}")
    if isinstance(v, int | float):
        return Decimal(str(v))
    texto = _texto(v)
    if "," in texto:  # 1.234,56
        texto = texto.replace(".", "").replace(",", ".")
    elif all(len(g) == 3 for g in texto.split(".")[1:]):  # 1.000 (milhar) ou 1000
        texto = texto.replace(".", "")
    elif texto.count(".") > 1:  # 1.2.3: nem milhar nem decimal
        raise ImportacaoError(f"{contexto}: número ilegível {v!r}")
    try:  # "0.5": um ponto sem três casas depois é o decimal
        return Decimal(texto)
    except InvalidOperation as erro:
        raise ImportacaoError(f"{contexto}: número ilegível {v!r}") from erro


def _codigo(produto: str) -> str:
    """'PETR4 - PETROBRAS' -> PETR4; 'Opção de Compra - PETRK300' -> PETRK300;
    'Futuro - WINZ25' -> WINZ25; 'Termo de Ação PETR4T' -> PETR4T."""
    partes = [p.strip() for p in produto.split(" - ")]
    if len(partes) > 1 and partes[0].startswith(("Opção de", "Futuro")):
        return partes[1].split()[0] if partes[1] else ""
    if produto.startswith("Termo de Ação "):
        return produto.removeprefix("Termo de Ação ").split()[0]
    return partes[0].split()[0] if partes[0] else ""


def movimentacao(caminho: Path) -> list[LinhaMovimentacao]:
    saida: list[LinhaMovimentacao] = []
    for i, x in enumerate(_linhas(caminho, ABA_MOVIMENTACAO,
                                  COLUNAS[TipoPlanilha.MOVIMENTACAO]), 2):  # fmt: skip
        contexto = f"{caminho.name}, linha {i}"
        sentido = _texto(x["Entrada/Saída"])
        if sentido not in ("Credito", "Debito"):
            raise ImportacaoError(f"{contexto}: Entrada/Saída {sentido!r}")
        produto = _texto(x["Produto"])
        saida.append(LinhaMovimentacao(
            data=_data(x["Data"], contexto), entrada=sentido == "Credito",
            movimentacao=_texto(x["Movimentação"]), produto=produto, codigo=_codigo(produto),
            instituicao=_texto(x["Instituição"]),
            quantidade=_numero(x["Quantidade"], contexto) or Decimal(0),
            preco=_numero(x["Preço unitário"], contexto),
            valor=_numero(x["Valor da Operação"], contexto),
        ))  # fmt: skip
    return saida


def negociacao(caminho: Path) -> list[LinhaNegociacao]:
    saida: list[LinhaNegociacao] = []
    for i, x in enumerate(_linhas(caminho, ABA_NEGOCIACAO,
                                  COLUNAS[TipoPlanilha.NEGOCIACAO]), 2):  # fmt: skip
        contexto = f"{caminho.name}, linha {i}"
        cv = {"Compra": "C", "Venda": "V"}.get(_texto(x["Tipo de Movimentação"]))
        if cv is None:
            raise ImportacaoError(f"{contexto}: Tipo de Movimentação {x['Tipo de Movimentação']!r}")
        quantidade, preco, valor = (_numero(x[c], contexto) for c in ("Quantidade", "Preço",
                                                                        "Valor"))  # fmt: skip
        if quantidade is None or preco is None or valor is None:
            raise ImportacaoError(f"{contexto}: negócio sem quantidade, preço ou valor")
        saida.append(LinhaNegociacao(
            data=_data(x["Data do Negócio"], contexto), cv=cv,  # type: ignore[arg-type]
            mercado=_texto(x["Mercado"]),
            vencimento=None if _vazio(x["Prazo/Vencimento"])
            else _data(x["Prazo/Vencimento"], contexto),
            instituicao=_texto(x["Instituição"]), codigo=_texto(x["Código de Negociação"]),
            quantidade=quantidade, preco=preco, valor=valor,
        ))  # fmt: skip
    return saida


def data_do_consolidado(caminho: Path) -> date:
    """Data da posição pelo nome do arquivo da B3: relatorio-consolidado-anual-AAAA -> 31/12/AAAA;
    relatorio-consolidado-mensal-AAAA-<mês> -> último dia do mês."""
    nome = caminho.stem.lower().replace("ç", "c")
    partes = nome.split("-")
    try:
        if "anual" in partes:
            return date(int(partes[partes.index("anual") + 1]), 12, 31)
        if "mensal" in partes:
            i = partes.index("mensal")
            ano, mes = int(partes[i + 1]), MESES.index(partes[i + 2]) + 1
            return date(ano + mes // 12, mes % 12 + 1, 1) - (date(2000, 1, 2) - date(2000, 1, 1))
    except (IndexError, ValueError) as erro:
        raise ImportacaoError(
            f"{caminho.name}: não deu para saber a data da posição pelo nome "
            "(relatorio-consolidado-anual-AAAA ou -mensal-AAAA-mês)"
        ) from erro
    raise ImportacaoError(f"{caminho.name}: nome fora do padrão do relatório consolidado da B3")


def posicao(caminho: Path, data: date | None = None) -> list[LinhaPosicao]:
    """Posição do consolidado (anual ou mensal) na `data` (padrão: pelo nome do arquivo): ações
    (inclui FII, ETF, BDR), ações emprestadas (doador soma, tomador subtrai: continuam do titular
    para o IR), opções (lançada negativa) e futuros (vendido negativo)."""
    data = data or data_do_consolidado(caminho)
    abas = set(_abas(caminho))
    saida: list[LinhaPosicao] = []
    acoes = _linhas(caminho, ABA_POSICAO_ACOES, COLUNAS_ACOES) if ABA_POSICAO_ACOES in abas else ()
    for i, x in enumerate(acoes, 2):
        contexto = f"{caminho.name}, ações, linha {i}"
        saida.append(LinhaPosicao(
            data=data, instituicao=_texto(x["Instituição"]), conta=_texto(x["Conta"]),
            codigo=_texto(x["Código de Negociação"]),
            isin=_texto(x["Código ISIN / Distribuição"]),
            quantidade=_numero(x["Quantidade"], contexto) or Decimal(0),
        ))  # fmt: skip
    if ABA_POSICAO_EMPRESTIMOS in abas:
        for i, x in enumerate(_linhas(caminho, ABA_POSICAO_EMPRESTIMOS, COLUNAS_EMPRESTIMOS), 2):
            contexto = f"{caminho.name}, empréstimos, linha {i}"
            natureza = _texto(x["Natureza"])
            if natureza not in ("Doador", "Tomador"):
                raise ImportacaoError(f"{contexto}: natureza {natureza!r}")
            quantidade = _numero(x["Quantidade"], contexto) or Decimal(0)
            saida.append(LinhaPosicao(
                data=data, instituicao=_texto(x["Instituição"]), conta="",
                codigo=_codigo(_texto(x["Produto"])), isin="",
                quantidade=quantidade if natureza == "Doador" else -quantidade,
                origem="emprestimo",
            ))  # fmt: skip
    derivativos: tuple[tuple[str, tuple[str, ...], Literal["opcao", "futuro"]], ...] = (
        (ABA_POSICAO_OPCOES, COLUNAS_OPCOES, "opcao"),
        (ABA_POSICAO_FUTURO, COLUNAS_FUTURO, "futuro"),
    )
    for aba, colunas, origem in derivativos:
        if aba not in abas:
            continue
        nomeadas: list[LinhaPosicao] = []
        sem_nome: list[LinhaPosicao] = []
        for i, x in enumerate(_linhas(caminho, aba, colunas), 2):
            contexto = f"{caminho.name}, {origem}, linha {i}"
            quantidade = _numero(x["Quantidade"], contexto) or Decimal(0)
            condicao = _texto(x["Condição"]).upper()  # "Lançador" ou "LANÇADOR", conforme o ano
            if condicao not in ("LANÇADOR", "COMPRADOR", "VENDIDO", "COMPRADO", "VENDEDOR"):
                raise ImportacaoError(f"{contexto}: condição {x['Condição']!r}")
            vendida = condicao in ("LANÇADOR", "VENDIDO", "VENDEDOR")
            produto = _texto(x["Produto"])
            codigo = _texto(x.get("Código de Negociação")) or _codigo(produto)
            destino = sem_nome if produto.startswith("-") else nomeadas
            destino.append(LinhaPosicao(
                data=data, instituicao=_texto(x["Instituição"]), conta="", codigo=codigo,
                isin="", quantidade=-quantidade if vendida else quantidade, origem=origem,
            ))  # fmt: skip
        saida += nomeadas + _sem_as_copias(sem_nome, nomeadas)
    return saida


def _sem_as_copias(
    sem_nome: list[LinhaPosicao], nomeadas: list[LinhaPosicao]
) -> list[LinhaPosicao]:
    """O consolidado mensal da B3 (visto em set/2026) repete séries com Produto " - Opção", sem o
    nome: a linha igual a uma com nome (instituição, código, quantidade) é cópia e sai; a sem
    par fica (é a única informação da posição)."""
    pares = Counter((x.instituicao, x.codigo, x.quantidade) for x in nomeadas)
    ficam: list[LinhaPosicao] = []
    for x in sem_nome:
        chave = (x.instituicao, x.codigo, x.quantidade)
        if pares[chave]:
            pares[chave] -= 1
        else:
            ficam.append(x)
    return ficam
