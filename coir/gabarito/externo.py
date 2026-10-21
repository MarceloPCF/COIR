"""Planilha externa de operações finalizadas (xlsx exportado por outro sistema de IR):
gabarito independente das vendas fechadas, venda a venda, para comparar com os movimentos do motor.

Layout (uma aba): seções "AAAA - Mês - Bovespa|BMF", cada uma com o cabeçalho (Ativo, Quantidade,
Corretora, C/V, Custo Médio, C/V, Custo Final, Data, Lucratividade, Observação, LAIR), as vendas e
uma linha "Total:"; no fim, o total geral. Observação: D = day trade, E = exercício, IS = isento,
DIS = day trade isento. O cabeçalho do arquivo traz nome e CPF do titular: não é lido.
"""

import csv
import re
from collections import defaultdict
from collections.abc import Callable, Iterable, Sequence
from dataclasses import dataclass, replace
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Literal

from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict

from coir.domain.fiscal import Movimento
from coir.domain.referencias import codigo_atual
from coir.erros import GabaritoError
from coir.gabarito.normalizar import corretora as normalizar_corretora
from coir.gabarito.normalizar import exigir_decimal

RE_SECAO = re.compile(r"^(?P<ano>\d{4}) - (?P<mes>[^-]+) - (?P<mercado>Bovespa|BMF)$")
ALIAS_CORRETORA = {"BANCO BTG PACTUAL S/A": "BTG"}
COLUNAS = ("Ativo", "Quantidade", "Corretora", "C/V", "Custo Médio", "C/V", "Custo Final", "Data",
           "Lucratividade", "Observação", "LAIR")  # fmt: skip
CENTAVO = Decimal("0.01")

Classe = Literal[
    "confere", "ticker com outro código", "outro dia do mesmo mês", "diferença",
    "só na planilha", "só no motor", "explicada",
]  # fmt: skip


class VendaExterna(BaseModel):
    model_config = ConfigDict(frozen=True, extra="forbid")

    data: date
    mercado: Literal["Bovespa", "BMF"]
    ticker: str
    quantidade: Decimal
    corretora: str
    abertura: str  # C (comprado), V (vendido) ou "-" (sem abertura: direito que expirou)
    custo_medio: Decimal
    fechamento: str  # "-" = vencimento / pó
    custo_final: Decimal
    observacao: str
    lair: Decimal  # lucro antes do IR

    @property
    def day_trade(self) -> bool:
        return self.observacao.startswith("D")


def ler_planilha_externa(arquivo: Path) -> list[VendaExterna]:
    """Lê as vendas; confere o total de cada seção e o total geral (divergência = erro)."""
    if not arquivo.is_file():
        raise GabaritoError(f"Planilha externa não encontrada: {arquivo}")
    livro = load_workbook(arquivo, read_only=False, data_only=True)
    linhas = [list(r) for r in livro.worksheets[0].iter_rows(values_only=True)]
    vendas: list[VendaExterna] = []
    secao: tuple[str, list[VendaExterna]] | None = None
    colunas: dict[int, int] | None = None
    secoes = 0
    for linha in linhas:
        textos = [c for c in linha if c is not None]
        if not textos:
            continue
        primeiro = str(textos[0]).strip()
        if RE_SECAO.match(primeiro):
            secao, colunas = (primeiro, []), None
            secoes += 1
            continue
        if secao is None:
            continue  # título e cabeçalho do arquivo (nome e CPF): ignorados
        if primeiro == "Ativo":
            colunas = _colunas(linha)
            continue
        if primeiro == "Total:":
            total = exigir_decimal(textos[-1], "Total")
            soma = sum((v.lair for v in secao[1]), Decimal(0))
            if secao[1] and abs(soma - total) >= CENTAVO:
                raise GabaritoError(f"Planilha externa, {secao[0]}: soma {soma} != total {total}")
            if not secao[1]:  # total geral
                soma = sum((v.lair for v in vendas), Decimal(0))
                if abs(soma - total) >= CENTAVO:
                    raise GabaritoError(f"Planilha externa: soma geral {soma} != total {total}")
            secao = (secao[0], [])
            continue
        if colunas is None:
            raise GabaritoError(f"Planilha externa, {secao[0]}: venda antes do cabeçalho")
        venda = _venda(linha, colunas, secao[0])
        secao[1].append(venda)
        vendas.append(venda)
    if not secoes:
        raise GabaritoError(f"{arquivo.name}: nenhuma seção 'AAAA - Mês - Bovespa|BMF'")
    return vendas


def _colunas(cabecalho: Sequence[object]) -> dict[int, int]:
    nomes = [(i, str(c).strip()) for i, c in enumerate(cabecalho) if c is not None]
    if tuple(n for _, n in nomes) != COLUNAS:
        raise GabaritoError(f"Planilha externa: cabeçalho inesperado {[n for _, n in nomes]}")
    return {k: i for k, (i, _) in enumerate(nomes)}


def _venda(linha: Sequence[object], col: dict[int, int], secao: str) -> VendaExterna:
    def c(k: int) -> object:
        return linha[col[k]]

    m = RE_SECAO.match(secao)
    if m is None:
        raise GabaritoError(f"Planilha externa: seção inválida {secao!r}")
    try:
        dia = datetime.strptime(str(c(7)), "%d/%m/%Y").date()
    except ValueError as erro:
        raise GabaritoError(f"Planilha externa, {secao}: data inválida {c(7)!r}") from erro
    nome = str(c(2)).strip().upper()
    return VendaExterna(
        data=dia, mercado=m["mercado"], ticker=str(c(0)).strip(),  # type: ignore[arg-type]
        quantidade=exigir_decimal(c(1), "Quantidade"),
        corretora=ALIAS_CORRETORA.get(nome) or normalizar_corretora(nome),
        abertura=str(c(3)).strip(), custo_medio=exigir_decimal(c(4), "Custo Médio"),
        fechamento=str(c(5)).strip(), custo_final=exigir_decimal(c(6), "Custo Final"),
        observacao=str(c(9) or "").strip(), lair=exigir_decimal(c(10), "LAIR"),
    )  # fmt: skip


# ------------------------------------------------------------------------------------ comparação
ticker_atual = codigo_atual  # nome antigo (a função mora no domínio)


@dataclass(frozen=True)
class LinhaComparacao:
    data: date
    corretora: str
    ticker: str
    motor: Decimal
    externo: Decimal
    classe: Classe
    motivo: str = ""

    @property
    def diferenca(self) -> Decimal:
        return self.motor - self.externo


@dataclass(frozen=True)
class MesComparado:
    ano: int
    mes: int
    mercado: str
    motor: Decimal
    externo: Decimal


Chave = tuple[date, str, str]


def comparar_externo(
    vendas: Iterable[VendaExterna], movimentos: Iterable[Movimento],
    ticker: Callable[[str], str], tolerancia: Decimal, mes: Decimal,
    desde: date | None = None,
) -> tuple[list[LinhaComparacao], list[MesComparado]]:  # fmt: skip
    """Resultado por (dia, corretora, ticker) e por mês e mercado. `ticker` normaliza os dois
    lados (trocas de código). Classes: ticker com outro código (mesmo dia, corretora e valor);
    outro dia do mesmo mês (mesma corretora e ticker, mesmo total no mês). `desde`: marco a
    partir do qual os dois lados são comparados (ex.: carteira zerada; antes dele a planilha pode
    ter começado com posições lançadas à mão)."""
    vendas = [v for v in vendas if desde is None or v.data >= desde]
    inicio = max(min((v.data for v in vendas), default=date.max), desde or date.min)
    externo: dict[Chave, Decimal] = defaultdict(Decimal)
    motor: dict[Chave, Decimal] = defaultdict(Decimal)
    meses: dict[tuple[int, int, str], list[Decimal]] = defaultdict(lambda: [Decimal(0)] * 2)
    for v in vendas:
        externo[(v.data, v.corretora, ticker(v.ticker))] += v.lair
        meses[(v.data.year, v.data.month, v.mercado)][1] += v.lair
    for m in movimentos:
        if m.data < inicio or not m.resultado:
            continue
        motor[(m.data, m.corretora or "?", ticker(m.ticker))] += m.resultado
        mercado = "BMF" if m.ativo is None else "Bovespa"
        meses[(m.data.year, m.data.month, mercado)][0] += m.resultado
    linhas = _classificar(motor, externo, tolerancia)
    comparados = [MesComparado(a, n, merc, *v) for (a, n, merc), v in sorted(meses.items())
                  if abs(v[0] - v[1]) > mes]  # fmt: skip
    return linhas, comparados


def carregar_explicacoes(arquivo: Path) -> dict[Chave, str]:
    """Divergências já analisadas (CSV em COIR_DADOS: data, corretora, ticker, motivo)."""
    if not arquivo.is_file():
        return {}
    with arquivo.open(encoding="utf-8-sig", newline="") as f:
        linhas = list(csv.DictReader(f))
    return {
        (date.fromisoformat(r["data"]), r["corretora"].strip(), r["ticker"].strip()): r["motivo"]
        for r in linhas
    }


def explicar(linhas: Iterable[LinhaComparacao], explicacoes: dict[Chave, str]
             ) -> list[LinhaComparacao]:  # fmt: skip
    """Marca como explicada a divergência com motivo registrado (a que já confere não muda)."""
    saida: list[LinhaComparacao] = []
    for x in linhas:
        motivo = explicacoes.get((x.data, x.corretora, x.ticker))
        if motivo and x.classe != "confere":
            x = replace(x, classe="explicada", motivo=motivo)
        saida.append(x)
    return saida


def _classificar(
    motor: dict[Chave, Decimal], externo: dict[Chave, Decimal], tol: Decimal
) -> list[LinhaComparacao]:
    classes: dict[Chave, Classe] = {}
    for k in set(motor) | set(externo):
        if abs(motor.get(k, Decimal(0)) - externo.get(k, Decimal(0))) <= tol:
            classes[k] = "confere"
        elif k not in externo:
            classes[k] = "só no motor"
        elif k not in motor:
            classes[k] = "só na planilha"
        else:
            classes[k] = "diferença"
    so_motor = [k for k, c in classes.items() if c == "só no motor"]
    for k in so_motor:  # mesmo dia, corretora e valor com outro código (série renomeada)
        par = next((s for s, c in classes.items() if c == "só na planilha" and s[:2] == k[:2]
                    and abs(externo[s] - motor[k]) <= tol), None)  # fmt: skip
        if par is not None:
            classes[k] = classes[par] = "ticker com outro código"
    abertas = [k for k, c in classes.items() if c in ("diferença", "só no motor", "só na planilha")]
    por_mes: dict[tuple[int, int, str, str], list[Chave]] = defaultdict(list)
    for k in abertas:
        por_mes[(k[0].year, k[0].month, k[1], k[2])].append(k)
    for ks in por_mes.values():
        if len(ks) > 1 and abs(sum(motor.get(k, Decimal(0)) - externo.get(k, Decimal(0))
                                   for k in ks)) <= tol:  # fmt: skip
            for k in ks:
                classes[k] = "outro dia do mesmo mês"
    return [
        LinhaComparacao(k[0], k[1], k[2], motor.get(k, Decimal(0)), externo.get(k, Decimal(0)), c)
        for k, c in sorted(classes.items())
    ]
