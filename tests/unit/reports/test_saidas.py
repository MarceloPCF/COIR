"""Escritores de relatório (ADR 0003): xlsx relido com openpyxl, csv pt-BR e PDF lido com
pypdfium2. Nenhum deles precisa do Excel."""

import csv
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import pypdfium2 as pdfium
import pytest
from openpyxl import load_workbook

from coir.erros import ExportacaoError
from coir.reports.modelo import (
    Cabecalho,
    Coluna,
    Relatorio,
    Secao,
    TipoColuna,
    cpf_formatado,
    formatar,
)
from coir.reports.saida_csv import escrever_csv
from coir.reports.saida_pdf import escrever_pdf
from coir.reports.saida_xlsx import escrever_xlsx
from tests.apoio import cpf_de_teste

D = Decimal
T = TipoColuna
AGORA = datetime(2026, 10, 16, 9, 30)
COLUNAS = (
    Coluna("Data", T.DATA), Coluna("Ativo"), Coluna("Quantidade", T.QUANTIDADE),
    Coluna("PM", T.PM), Coluna("Resultado", T.DINHEIRO), Coluna("%", T.PERCENTUAL),
)  # fmt: skip


def _relatorio(linhas: int = 3) -> Relatorio:
    corpo = tuple(
        (date(2025, 1, 2 + i % 20), f"AAAA{i % 9}", D("100.5"), D("12.12345678"),
         D("-1234.567") if i % 2 else D("1000"), D("0.1234"))
        for i in range(linhas)
    )  # fmt: skip
    total = ("Total", None, None, None, sum((x[4] for x in corpo), D(0)), None)
    return Relatorio(
        chave="teste", titulo="Operações finalizadas",
        cabecalho=Cabecalho("Titular de Teste", cpf_de_teste(), "Ano-calendário 2025", AGORA),
        secoes=(Secao("2025 - Janeiro - Bovespa", COLUNAS, corpo, total, nota="Lei 11.033/2004"),
                Secao("2025 - Fevereiro - Bovespa", COLUNAS, corpo[:1])),
        pdf=True,
    )  # fmt: skip


def test_formatacao_pt_br_por_tipo() -> None:
    assert formatar(D("1234567.891"), T.DINHEIRO) == "1.234.567,89"
    assert formatar(D("1234.5"), T.DINHEIRO, milhar=False) == "1234,50"
    assert formatar(D("12.123456789"), T.PM) == "12,12345679"
    assert formatar(D("100.50"), T.QUANTIDADE) == "100,5"
    assert formatar(D(2000), T.QUANTIDADE) == "2.000"
    assert formatar(D("0.1234"), T.PERCENTUAL) == "12,34%"
    assert formatar(date(2025, 3, 7), T.DATA) == "07/03/2025" and formatar(None, T.TEXTO) == ""
    c = cpf_de_teste()
    assert cpf_formatado(c) == f"{c[:3]}.{c[3:6]}.{c[6:9]}-{c[9:]}"


def test_linha_com_largura_errada_falha_alto() -> None:
    with pytest.raises(ExportacaoError, match="2 valores para 6 colunas"):
        Secao("x", COLUNAS, (("a", "b"),))


def test_xlsx_so_valores_com_formatos_e_cabecalho(tmp_path: Path) -> None:
    rel = _relatorio()
    caminho = escrever_xlsx([rel, rel], tmp_path / "r.xlsx", AGORA)
    livro = load_workbook(caminho)
    assert livro.sheetnames == ["Operações finalizadas", "Operações finalizadas (2)"]
    aba = livro.worksheets[0]
    assert aba["A1"].value == "Operações finalizadas"
    assert aba["A2"].value == f"Titular de Teste - CPF {cpf_formatado(cpf_de_teste())}"
    assert aba["A6"].value == "Data" and aba["D6"].value == "PM"
    assert aba["A7"].value == datetime(2025, 1, 2) and aba["A7"].number_format == "dd/mm/yyyy"
    assert aba["D7"].value == pytest.approx(12.12345678) and aba["D7"].number_format.startswith(
        "#,##0.00000000")  # fmt: skip
    assert aba["E8"].value == pytest.approx(-1234.57)
    assert aba["A10"].value == "Total" and aba["E10"].value == pytest.approx(765.43)
    formulas = [c for linha in aba.iter_rows() for c in linha if c.data_type == "f"]
    assert formulas == []
    assert livro.properties.created == AGORA


def test_csv_padrao_brasileiro_com_secoes(tmp_path: Path) -> None:
    [caminho] = escrever_csv(_relatorio(), tmp_path, "coir-2025-titular1")
    assert caminho.name == "coir-2025-titular1-teste.csv"
    assert caminho.read_bytes().startswith(b"\xef\xbb\xbf")
    with caminho.open(encoding="utf-8-sig", newline="") as f:
        linhas = list(csv.reader(f, delimiter=";"))
    assert linhas[0] == ["Seção", "Data", "Ativo", "Quantidade", "PM", "Resultado", "%"]
    assert linhas[1] == ["2025 - Janeiro - Bovespa", "02/01/2025", "AAAA0", "100,5", "12,12345678",
                         "1000,00", "12,34%"]  # fmt: skip
    assert linhas[4][1:] == ["Total", "", "", "", "765,43", ""]
    assert linhas[5][0] == "2025 - Fevereiro - Bovespa" and len(linhas) == 6


def test_csv_separa_secoes_com_colunas_diferentes(tmp_path: Path) -> None:
    outra = Secao("Resumo do ano", (Coluna("Mês"), Coluna("DARF", T.DINHEIRO)), (("jan", D(10)),))
    rel = _relatorio()
    rel = Relatorio(rel.chave, rel.titulo, rel.cabecalho, (*rel.secoes, outra))
    nomes = sorted(p.name for p in escrever_csv(rel, tmp_path, "p"))
    assert nomes == ["p-teste-2025-janeiro-bovespa.csv", "p-teste-resumo-do-ano.csv"]


def test_pdf_paginas_com_cabecalho_rodape_e_total(tmp_path: Path) -> None:
    caminho = escrever_pdf(_relatorio(linhas=120), tmp_path / "r.pdf")
    assert caminho.read_bytes().startswith(b"%PDF")
    pdf = pdfium.PdfDocument(caminho)
    paginas = len(pdf)
    assert paginas >= 2  # tabela longa quebra de página
    textos = [pdf[i].get_textpage().get_text_range() for i in range(paginas)]
    assert all(f"Página {i + 1} de {paginas}" in t for i, t in enumerate(textos))
    assert all(cpf_formatado(cpf_de_teste()) in t and "Operações finalizadas" in t for t in textos)
    assert all("Data" in t and "Resultado" in t for t in textos)  # cabeçalho repetido
    assert "Total" in "".join(textos) and "não substitui" in textos[-1]
    pdf.close()


def test_pdf_reproduzivel(tmp_path: Path) -> None:
    a = escrever_pdf(_relatorio(), tmp_path / "a.pdf").read_bytes()
    b = escrever_pdf(_relatorio(), tmp_path / "b.pdf").read_bytes()
    assert a == b
