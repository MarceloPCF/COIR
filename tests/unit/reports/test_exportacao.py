"""Exportação de ponta a ponta (F5): banco das notas de fixture -> relatórios -> xlsx, csv e pdf,
pela função de serviço e pelo comando `coir exportar`. Nada precisa do Excel."""

from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook
from typer.testing import CliRunner

from coir import referencias
from coir.cli import app
from coir.config import RAIZ_REPOSITORIO
from coir.db import repositorio as repo
from coir.db.sessao import abrir, migrar
from coir.erros import ExportacaoError
from coir.servicos.importacao import importar_arquivos
from coir.servicos.relatorios import dados_banco, exportar
from tests.apoio import cpf_de_teste
from tests.unit.db.test_banco import FIXTURES, extrator_txt

AGORA = datetime(2026, 10, 16, 10, 0)
REFS = referencias.carregar()
FISCAIS = referencias.carregar_fiscais()


@pytest.fixture
def banco(tmp_path: Path) -> Path:
    caminho = tmp_path / "dados" / "coir.db"
    migrar(caminho, AGORA)
    with abrir(caminho).begin() as s:
        repo.adicionar_titular(s, cpf_de_teste(), "Titular de Teste")
    pasta = tmp_path / "notas"
    pasta.mkdir()
    for f in FIXTURES.glob("*.txt"):
        if "exercicio" not in f.stem:  # a opção exercida não está nas fixtures
            (pasta / f.name).write_bytes(f.read_bytes())
    with abrir(caminho)() as s:
        importar_arquivos(s, sorted(pasta.iterdir()), REFS, AGORA, extrator=extrator_txt)
    return caminho


def test_exporta_todos_os_formatos_sem_cpf_no_nome(banco: Path, tmp_path: Path) -> None:
    dados = dados_banco(banco, REFS, FISCAIS, 2023, AGORA)
    assert dados.cpf == cpf_de_teste() and dados.notas  # notas de 2023 do titular
    assert all(n.nota.data_pregao.year == 2023 for n in dados.notas)
    arquivos = exportar(dados, tmp_path / "saida")
    nomes = {a.name for a in arquivos}
    assert "coir-2023-titular1.xlsx" in nomes
    assert "coir-2023-titular1-resumo-fiscal.pdf" in nomes
    assert any(n.endswith(".csv") for n in nomes)
    assert not any(cpf_de_teste() in n for n in nomes)
    livro = load_workbook(tmp_path / "saida" / "coir-2023-titular1.xlsx")
    assert "Livro de notas" in livro.sheetnames and "Resumo fiscal do ano" in livro.sheetnames


def test_destino_no_repositorio_e_opcoes_invalidas_falham(banco: Path, tmp_path: Path) -> None:
    dados = dados_banco(banco, REFS, FISCAIS, 2023, AGORA)
    with pytest.raises(ExportacaoError, match="repositório"):
        exportar(dados, RAIZ_REPOSITORIO / "saida-teste")
    with pytest.raises(ExportacaoError, match="formato"):
        exportar(dados, tmp_path / "x", ("docx",))
    with pytest.raises(ExportacaoError, match="relatório desconhecido"):
        exportar(dados, tmp_path / "x", chaves=("nada",))
    with pytest.raises(ExportacaoError, match="titular 9"):
        dados_banco(banco, REFS, FISCAIS, 2023, AGORA, titular=9)
    with pytest.raises(ExportacaoError, match="ainda não começou"):
        dados_banco(banco, REFS, FISCAIS, 2027, AGORA)


def test_ano_em_curso_apura_so_ate_hoje(banco: Path) -> None:
    """Uma opção que vence depois da data de geração não vira pó no relatório do ano em curso."""
    dados = dados_banco(banco, REFS, FISCAIS, 2026, datetime(2026, 8, 31, 12, 0))
    assert max(m.data for m in dados.apuracao.movimentos) <= datetime(2026, 8, 31).date()
    assert dados.data_carteira == datetime(2026, 8, 31).date()


def test_comando_exportar(banco: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("COIR_DADOS", str(banco.parent))
    destino = tmp_path / "rel"
    args = ["exportar", "--ano", "2023", "--formato", "csv,xlsx", "--relatorio", "darfs,carteira"]
    res = CliRunner().invoke(app, [*args, "--destino", str(destino)])
    assert res.exit_code == 0, res.output
    assert "coir-2023-titular1.xlsx" in res.output and "DARF do ano" in res.output
    assert sorted(p.suffix for p in destino.iterdir()) == [".csv"] * 3 + [".xlsx"]
