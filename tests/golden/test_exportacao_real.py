"""F5: relatórios de todos os anos das notas reais (+ eventos.csv), sem Excel. O DARF do ano no
xlsx é o da apuração. O banco e as saídas são temporários (nunca a pasta do usuário)."""

import time
from datetime import datetime
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from coir import referencias
from coir.config import dados_dir
from coir.db import repositorio as repo
from coir.db.sessao import abrir, migrar
from coir.servicos.eventos import importar_eventos, ler_csv
from coir.servicos.importacao import RelatorioImportacao, arquivos_da_pasta, importar_arquivos
from coir.servicos.relatorios import dados_banco, exportar
from tests.apoio import cpf_de_teste

pytestmark = pytest.mark.golden
DADOS = dados_dir()
AGORA = datetime(2026, 10, 16, 12, 0, 0)


@pytest.fixture(scope="module")
def banco(tmp_path_factory: pytest.TempPathFactory) -> Path:
    notas, eventos = DADOS / "_notas-corretagens", DADOS / "eventos.csv"
    if not (notas.is_dir() and eventos.is_file()):
        pytest.skip("notas ou eventos.csv ausentes")
    caminho = tmp_path_factory.mktemp("banco") / "coir.db"
    migrar(caminho, AGORA)
    fabrica = abrir(caminho)
    with fabrica.begin() as s:
        repo.adicionar_titular(s, cpf_de_teste(), "Titular de Teste")
    refs = referencias.carregar()
    with fabrica() as s:
        importar_arquivos(s, arquivos_da_pasta(notas, RelatorioImportacao()), refs, AGORA)
    with fabrica() as s:
        importar_eventos(s, ler_csv(eventos))
    fabrica.kw["bind"].dispose()  # type: ignore[union-attr]
    return caminho


def test_todos_os_anos_exportam_e_o_darf_confere(banco: Path, tmp_path: Path) -> None:
    refs, fiscais = referencias.carregar(), referencias.carregar_fiscais()
    for ano in range(2008, 2027):
        inicio = time.perf_counter()
        dados = dados_banco(banco, refs, fiscais, ano, AGORA)
        formatos = ("xlsx", "csv", "pdf") if ano == 2021 else ("xlsx",)
        exportar(dados, tmp_path / str(ano), formatos)
        segundos = time.perf_counter() - inicio
        assert segundos < 30, f"{ano}: {segundos:.1f} s"
        livro = load_workbook(tmp_path / str(ano) / f"coir-{ano}-titular1.xlsx")
        aba = livro["Resumo fiscal do ano"]
        linhas = list(aba.iter_rows(values_only=True))
        cab = next(i for i, x in enumerate(linhas) if x[0] == "Mês" and "DARF" in x)
        col = linhas[cab].index("DARF")
        total = next(x for x in linhas[cab + 1 :] if x[0] == "Total")
        darf = sum((m.darf for m in dados.apuracao.meses if m.ano == ano), Decimal(0))
        assert Decimal(str(total[col])).quantize(Decimal("0.01")) == darf, ano
